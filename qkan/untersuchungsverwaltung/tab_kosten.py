# untersuchungsverwaltung/tab_kosten.py

import os
import json
import subprocess

import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows

from qgis.core import QgsProject
from qgis.PyQt.QtWidgets import QMessageBox, QTableWidgetItem, QFileDialog


class KostenManager:
    """
    Verwaltet die Logik für den Tab 'Kostenermittlung'.

    Preislogik:
    - dimensionsabhängig: Reinigung, Reinigung TV, TV, Panoramo
    - Einzelpreis: GAL, Panoramo SI

    Die Summierung erfolgt weiterhin getrennt nach:
    - Mischwasser
    - Regenwasser
    - Schmutzwasser
    """

    # ==================================================================
    # INITIALISIERUNG
    # ==================================================================

    SINGLE_PRICE_KEY = "default"

    def __init__(self, dialog):
        self.dialog = dialog

        self.reinigung_costs = {}
        self.reinigung_TV_costs = {}
        self.TV_costs = {}
        self.Panoramo_costs = {}

        self.gal_single_price = 0.0
        self.panoramo_si_single_price = 0.0

        self.load_cleaning_costs()

    # ==================================================================
    # PREISDATEI / PREISLADEN
    # ==================================================================

    def _get_price_file_path(self):
        """Liefert den Pfad zur JSON-Preisliste im Tool-Ordner."""
        return os.path.join(
            os.path.dirname(__file__),
            "preisliste_untersuchung.json"
        )

    def _get_default_price_structure(self):
        """Standardstruktur für die Preisdatei."""
        return {
            "Reinigung": {},
            "Reinigung TV": {},
            "TV": {},
            "GAL": {self.SINGLE_PRICE_KEY: 0.0},
            "Panoramo": {},
            "Panoramo SI": {self.SINGLE_PRICE_KEY: 0.0},
        }

    def load_cleaning_costs(self):
        """
        Lädt die Preise aus der JSON-Datei.

        Erwartete Struktur:
        {
            "Reinigung": {"300": 1.50},
            "Reinigung TV": {"300": 2.10},
            "TV": {"300": 3.50},
            "GAL": {"default": 19.59},
            "Panoramo": {"300": 2.80},
            "Panoramo SI": {"default": 37.51}
        }
        """
        file_path = self._get_price_file_path()
        data = self._get_default_price_structure()

        if os.path.exists(file_path):
            with open(file_path, "r", encoding="utf-8") as file:
                try:
                    loaded = json.load(file)

                    for key in data.keys():
                        value = loaded.get(key, {})
                        if isinstance(value, dict):
                            data[key] = value

                except json.JSONDecodeError:
                    QMessageBox.warning(
                        self.dialog,
                        "Fehler",
                        "Fehler beim Laden der Kostenliste."
                    )

        self.reinigung_costs = data.get("Reinigung", {})
        self.reinigung_TV_costs = data.get("Reinigung TV", {})
        self.TV_costs = data.get("TV", {})
        self.Panoramo_costs = data.get("Panoramo", {})

        gal_data = data.get("GAL", {})
        if isinstance(gal_data, dict):
            self.gal_single_price = self._parse_float(
                gal_data.get(self.SINGLE_PRICE_KEY, 0.0),
                default=0.0
            )
        else:
            self.gal_single_price = 0.0

        panoramo_si_data = data.get("Panoramo SI", {})
        if isinstance(panoramo_si_data, dict):
            self.panoramo_si_single_price = self._parse_float(
                panoramo_si_data.get(self.SINGLE_PRICE_KEY, 0.0),
                default=0.0
            )
        else:
            self.panoramo_si_single_price = 0.0

    # ==================================================================
    # HILFSMETHODEN
    # ==================================================================

    def _normalize_dimension(self, dimension):
        """Normalisiert eine Dimension in einen vergleichbaren String."""
        return str(dimension).strip()

    def _get_dimension_price(self, price_dict, dimension):
        """Liefert den dimensionsabhängigen Preis für eine Kategorie."""
        return self._parse_float(
            price_dict.get(self._normalize_dimension(dimension)),
            default=None
        )

    def _parse_float(self, value, default=None):
        """
        Wandelt Werte robust in float um.
        Unterstützt z.B. '12,5' oder '12.5'.
        """
        if value is None:
            return default

        text = str(value).strip()
        if not text or text.lower() == "none":
            return default

        try:
            return float(text.replace(",", "."))
        except ValueError:
            return default

    def _map_system(self, system_text):
        """
        Ordnet verschiedene Entwässerungssystem-Bezeichnungen
        auf Mischwasser / Regenwasser / Schmutzwasser ab.
        """
        system_mapping = {
            "Mischwasser": ["Mischwasser", "KM Mischwasserkanal"],
            "Regenwasser": ["Regenwasser", "KR Regenwasserkanal", "BA Bachverrohrung"],
            "Schmutzwasser": ["Schmutzwasser", "KS Schmutzwasserkanal"],
        }

        for mapped_name, aliases in system_mapping.items():
            if system_text in aliases:
                return mapped_name

        return system_text

    def _set_currency_item(self, row, column, value):
        """Schreibt einen Euro-Wert formatiert in die Tabelle."""
        self.dialog.Liste_Haltungen.setItem(
            row,
            column,
            QTableWidgetItem(f"{value:.2f} €")
        )

    def _set_text_item(self, row, column, text):
        """Schreibt einen Textwert in die Tabelle."""
        self.dialog.Liste_Haltungen.setItem(
            row,
            column,
            QTableWidgetItem(str(text))
        )

    def _set_summary_fields(self, netto_mw, netto_rw, netto_sw, prefix):
        """
        Setzt Netto- und Brutto-Summenfelder anhand eines Prefixes.

        Beispiel prefix='Reinigung':
        - Reinigung_netto_MW
        - Reinigung_brutto_MW
        - Reinigung_netto_RW
        - Reinigung_brutto_RW
        - Reinigung_netto_SW
        - Reinigung_brutto_SW
        """
        mw_netto = getattr(self.dialog, f"{prefix}_netto_MW", None)
        mw_brutto = getattr(self.dialog, f"{prefix}_brutto_MW", None)
        rw_netto = getattr(self.dialog, f"{prefix}_netto_RW", None)
        rw_brutto = getattr(self.dialog, f"{prefix}_brutto_RW", None)
        sw_netto = getattr(self.dialog, f"{prefix}_netto_SW", None)
        sw_brutto = getattr(self.dialog, f"{prefix}_brutto_SW", None)

        if mw_netto:
            mw_netto.setText(f"{netto_mw:.2f} €")
        if mw_brutto:
            mw_brutto.setText(f"{netto_mw * 1.19:.2f} €")

        if rw_netto:
            rw_netto.setText(f"{netto_rw:.2f} €")
        if rw_brutto:
            rw_brutto.setText(f"{netto_rw * 1.19:.2f} €")

        if sw_netto:
            sw_netto.setText(f"{netto_sw:.2f} €")
        if sw_brutto:
            sw_brutto.setText(f"{netto_sw * 1.19:.2f} €")

    # ==================================================================
    # UI-REAKTIONEN
    # ==================================================================

    def on_checkbox_changed(self):
        """Aktualisiert die Berechnungen nach Änderung der Zusatzoptionen."""
        self.calculateCleaningCost()
        self.calculateCleaningTVCost()

    def showSelectedFeatures(self):
        """Lädt die aktuell selektierten Haltungen aus dem aktiven Layer."""
        active_layer = self.dialog.iface.activeLayer()
        if not active_layer:
            QMessageBox.warning(
                self.dialog,
                "Fehler",
                "Kein aktiver Layer gefunden."
            )
            return

        selected_features = active_layer.selectedFeatures()
        self.dialog.Liste_Haltungen.clearContents()
        self.dialog.Liste_Haltungen.setRowCount(0)
        self.populateTableWidget(selected_features)

    # ==================================================================
    # DATENBANKABFRAGEN
    # ==================================================================

    def get_gal_counts_from_db(self, haltnams):
        """Ermittelt optional die GAL-Anzahlen aus der Datenbank."""
        if not self.dialog.cur:
            return {}

        gal_counts = {}

        try:
            haltnams_list = list(haltnams)
            if not haltnams_list:
                return {}

            if self.dialog.is_spatialite:
                placeholders = ",".join(["?"] * len(haltnams_list))
                query = f"""
                    SELECT u.untersuchhal, COUNT(*) as gal_anzahl
                    FROM untersuchdat_haltung u
                    JOIN (
                        SELECT untersuchhal, MAX(untersuchtag) AS max_tag
                        FROM untersuchdat_haltung
                        WHERE untersuchhal IN ({placeholders})
                        GROUP BY untersuchhal
                    ) latest
                      ON u.untersuchhal = latest.untersuchhal
                     AND u.untersuchtag = latest.max_tag
                    WHERE u.kuerzel = ? AND u.charakt2 = ?
                    GROUP BY u.untersuchhal
                """
                params = haltnams_list + ["BCA", "A"]
                self.dialog.cur.execute(query, params)

            else:
                query = """
                    SELECT u.untersuchhal, COUNT(*) as gal_anzahl
                    FROM untersuchdat_haltung u
                    JOIN (
                        SELECT untersuchhal, MAX(untersuchtag) AS max_tag
                        FROM untersuchdat_haltung
                        WHERE untersuchhal = ANY(%s)
                        GROUP BY untersuchhal
                    ) latest
                      ON u.untersuchhal = latest.untersuchhal
                     AND u.untersuchtag = latest.max_tag
                    WHERE u.kuerzel = %s AND u.charakt2 = %s
                    GROUP BY u.untersuchhal
                """
                self.dialog.cur.execute(query, (haltnams_list, "BCA", "A"))

            rows = self.dialog.cur.fetchall()
            gal_counts = {row[0]: row[1] for row in rows}

        except Exception as e:
            self.dialog.conn.rollback()
            QMessageBox.warning(
                self.dialog,
                "Datenbankfehler",
                f"Fehler bei GAL-Abfrage: {e}"
            )

        return gal_counts

    # ==================================================================
    # TABELLENBEFÜLLUNG
    # ==================================================================

    def populateTableWidget(self, selected_features):
        """Befüllt die Haltungstabelle mit den selektierten Features."""
        if not self.dialog.cur:
            QMessageBox.warning(
                self.dialog,
                "Fehler",
                "Keine aktive Datenbankverbindung."
            )
            return

        active_layer = self.dialog.iface.activeLayer()
        if not active_layer or active_layer.name() != "Haltungen":
            msg_box = QMessageBox()
            msg_box.setIcon(QMessageBox.Question)
            msg_box.setWindowTitle("Layer-Auswahl")
            msg_box.setText(
                "Der Layer 'Haltungen' ist nicht aktiv. Soll er jetzt ausgewählt werden?"
            )
            msg_box.setStandardButtons(QMessageBox.Yes | QMessageBox.No)

            if msg_box.exec_() == QMessageBox.Yes:
                h_layer_list = QgsProject.instance().mapLayersByName("Haltungen")
                if h_layer_list:
                    h_layer = h_layer_list[0]
                    self.dialog.iface.setActiveLayer(h_layer)
                    selected_features = h_layer.selectedFeatures()
                    self.populateTableWidget(selected_features)
                    return

        selected_haltnams = {f["haltnam"] for f in selected_features}
        if not selected_haltnams:
            return

        try:
            selected_haltnams_list = list(selected_haltnams)

            if self.dialog.is_spatialite:
                placeholders = ",".join(["?"] * len(selected_haltnams_list))
                query = (
                    f"SELECT haltnam, laenge, entwart, hoehe, breite "
                    f"FROM haltungen WHERE haltnam IN ({placeholders})"
                )
                self.dialog.cur.execute(query, selected_haltnams_list)
            else:
                query = """
                    SELECT haltnam, laenge, entwart, hoehe, breite
                    FROM public.haltungen
                    WHERE haltnam = ANY(%s)
                """
                self.dialog.cur.execute(query, (selected_haltnams_list,))

            filtered_features = self.dialog.cur.fetchall()

        except Exception as e:
            self.dialog.conn.rollback()
            QMessageBox.warning(
                self.dialog,
                "Datenbankfehler",
                f"Fehler bei SQL-Abfrage: {e}"
            )
            return

        self.dialog.Liste_Haltungen.setRowCount(len(filtered_features))

        msg_box = QMessageBox()
        msg_box.setIcon(QMessageBox.Question)
        msg_box.setWindowTitle("Abfrage")
        msg_box.setText("Soll die Anzahl der GALs aus der Datenbank ermittelt werden?")
        msg_box.setStandardButtons(QMessageBox.Yes | QMessageBox.No)
        fetch_gal = msg_box.exec_() == QMessageBox.Yes

        gal_counts = {}
        if fetch_gal:
            gal_counts = self.get_gal_counts_from_db(selected_haltnams)

        self.dialog.Liste_Haltungen.setUpdatesEnabled(False)
        try:
            for i, feature in enumerate(filtered_features):
                haltnam, laenge, entwart, hoehe, breite = feature

                h = int(hoehe) if hoehe is not None else 0
                b = int(breite) if breite is not None else 0
                dimension = f"{b}/{h}" if h != b else str(h)

                anzahl_gal = (
                    gal_counts.get(haltnam, 0)
                    if fetch_gal
                    else "Nicht berechnet"
                )

                values = [
                    haltnam,
                    laenge,
                    entwart,
                    dimension,
                    anzahl_gal,
                    "",
                    "",
                    "",
                    "",
                    "",
                ]

                for j, value in enumerate(values):
                    item = QTableWidgetItem(str(value) if value is not None else "")
                    self.dialog.Liste_Haltungen.setItem(i, j, item)

        finally:
            self.dialog.Liste_Haltungen.setUpdatesEnabled(True)

    # ==================================================================
    # KOSTENBERECHNUNG: REINIGUNG
    # ==================================================================

    def calculateCleaningCost(self):
        """Berechnet die Reinigungskosten und summiert sie je System."""
        self.dialog.Liste_Haltungen.setUpdatesEnabled(False)

        try:
            total_cost_mw = 0.0
            total_cost_rw = 0.0
            total_cost_sw = 0.0

            for row in range(self.dialog.Liste_Haltungen.rowCount()):
                dimension_item = self.dialog.Liste_Haltungen.item(row, 3)
                length_item = self.dialog.Liste_Haltungen.item(row, 1)
                system_item = self.dialog.Liste_Haltungen.item(row, 2)

                if not (dimension_item and length_item and system_item):
                    continue

                dimension = self._normalize_dimension(dimension_item.text())
                length = self._parse_float(length_item.text(), default=None)
                mapped_system = self._map_system(system_item.text())

                if length is None:
                    continue

                cost_per_unit = self._get_dimension_price(
                    self.reinigung_costs,
                    dimension
                )

                if cost_per_unit is None:
                    self._set_text_item(row, 5, "kein Preis")
                    continue

                cost = length * cost_per_unit
                brutto_cost = cost * 1.19
                self._set_currency_item(row, 5, brutto_cost)

                if mapped_system == "Mischwasser":
                    total_cost_mw += cost
                elif mapped_system == "Regenwasser":
                    total_cost_rw += cost
                elif mapped_system == "Schmutzwasser":
                    total_cost_sw += cost

            self._set_summary_fields(
                netto_mw=total_cost_mw,
                netto_rw=total_cost_rw,
                netto_sw=total_cost_sw,
                prefix="Reinigung",
            )

        finally:
            self.dialog.Liste_Haltungen.setUpdatesEnabled(True)

    # ==================================================================
    # KOSTENBERECHNUNG: TV / GAL / PANORAMO
    # ==================================================================

    def calculateCleaningTVCost(self):
        """
        Berechnet:
        - Reinigung für Befahrung
        - TV
        - optional GAL (Einzelpreis, nicht dimensionsabhängig)
        - optional Panoramo
        - optional Panoramo SI (Einzelpreis je Haltung)

        Summierung erfolgt getrennt nach MW / RW / SW.
        """
        self.dialog.Liste_Haltungen.setUpdatesEnabled(False)

        try:
            total_cost_cleaning_tv_mw = 0.0
            total_cost_cleaning_tv_rw = 0.0
            total_cost_cleaning_tv_sw = 0.0

            total_cost_tv_mw = 0.0
            total_cost_tv_rw = 0.0
            total_cost_tv_sw = 0.0

            total_cost_gal_mw = 0.0
            total_cost_gal_rw = 0.0
            total_cost_gal_sw = 0.0

            total_cost_panoramo_mw = 0.0
            total_cost_panoramo_rw = 0.0
            total_cost_panoramo_sw = 0.0

            for row in range(self.dialog.Liste_Haltungen.rowCount()):
                dimension_item = self.dialog.Liste_Haltungen.item(row, 3)
                length_item = self.dialog.Liste_Haltungen.item(row, 1)
                system_item = self.dialog.Liste_Haltungen.item(row, 2)
                gal_item = self.dialog.Liste_Haltungen.item(row, 4)

                if not (dimension_item and length_item and system_item):
                    continue

                dimension = self._normalize_dimension(dimension_item.text())
                length = self._parse_float(length_item.text(), default=None)
                mapped_system = self._map_system(system_item.text())
                gal_value = self._parse_float(
                    gal_item.text() if gal_item else "0",
                    default=0.0
                )

                if length is None:
                    continue

                # ------------------------------------------------------
                # Reinigung TV (dimensionsabhängig)
                # ------------------------------------------------------
                cleaning_tv_unit = self._get_dimension_price(
                    self.reinigung_TV_costs,
                    dimension
                )
                if cleaning_tv_unit is not None:
                    cost = length * cleaning_tv_unit
                    self._set_currency_item(row, 6, cost * 1.19)

                    if mapped_system == "Mischwasser":
                        total_cost_cleaning_tv_mw += cost
                    elif mapped_system == "Regenwasser":
                        total_cost_cleaning_tv_rw += cost
                    elif mapped_system == "Schmutzwasser":
                        total_cost_cleaning_tv_sw += cost
                else:
                    self._set_text_item(row, 6, "kein Preis")

                # ------------------------------------------------------
                # TV (dimensionsabhängig)
                # ------------------------------------------------------
                tv_unit = self._get_dimension_price(self.TV_costs, dimension)
                if tv_unit is not None:
                    cost = length * tv_unit
                    self._set_currency_item(row, 7, cost * 1.19)

                    if mapped_system == "Mischwasser":
                        total_cost_tv_mw += cost
                    elif mapped_system == "Regenwasser":
                        total_cost_tv_rw += cost
                    elif mapped_system == "Schmutzwasser":
                        total_cost_tv_sw += cost
                else:
                    self._set_text_item(row, 7, "kein Preis")

                # ------------------------------------------------------
                # GAL (Einzelpreis, unabhängig von Dimension)
                # ------------------------------------------------------
                if self.dialog.checkBox_GAL.isChecked():
                    if self.gal_single_price is not None:
                        cost_gal = gal_value * self.gal_single_price
                        self._set_currency_item(row, 8, cost_gal * 1.19)

                        if mapped_system == "Mischwasser":
                            total_cost_gal_mw += cost_gal
                        elif mapped_system == "Regenwasser":
                            total_cost_gal_rw += cost_gal
                        elif mapped_system == "Schmutzwasser":
                            total_cost_gal_sw += cost_gal
                    else:
                        self._set_text_item(row, 8, "kein Preis")
                else:
                    self._set_text_item(row, 8, "")

                # ------------------------------------------------------
                # Panoramo (dimensionsabhängig + Panoramo SI Einzelpreis)
                # ------------------------------------------------------
                if self.dialog.checkBox_Panoramo.isChecked():
                    panoramo_unit = self._get_dimension_price(
                        self.Panoramo_costs,
                        dimension
                    )

                    if panoramo_unit is not None:
                        cost_panoramo = (length * panoramo_unit) + self.panoramo_si_single_price
                        self._set_currency_item(row, 9, cost_panoramo * 1.19)

                        if mapped_system == "Mischwasser":
                            total_cost_panoramo_mw += cost_panoramo
                        elif mapped_system == "Regenwasser":
                            total_cost_panoramo_rw += cost_panoramo
                        elif mapped_system == "Schmutzwasser":
                            total_cost_panoramo_sw += cost_panoramo
                    else:
                        self._set_text_item(row, 9, "kein Preis")
                else:
                    self._set_text_item(row, 9, "")

            # Summen Reinigung für Befahrung
            self._set_summary_fields(
                netto_mw=total_cost_cleaning_tv_mw,
                netto_rw=total_cost_cleaning_tv_rw,
                netto_sw=total_cost_cleaning_tv_sw,
                prefix="Reinigung_befahrung",
            )

            # Summen Befahrung gesamt
            total_cost_tv_mw_all = total_cost_tv_mw + total_cost_gal_mw + total_cost_panoramo_mw
            total_cost_tv_rw_all = total_cost_tv_rw + total_cost_gal_rw + total_cost_panoramo_rw
            total_cost_tv_sw_all = total_cost_tv_sw + total_cost_gal_sw + total_cost_panoramo_sw

            self._set_summary_fields(
                netto_mw=total_cost_tv_mw_all,
                netto_rw=total_cost_tv_rw_all,
                netto_sw=total_cost_tv_sw_all,
                prefix="Befahrung",
            )

        finally:
            self.dialog.Liste_Haltungen.setUpdatesEnabled(True)

    # ==================================================================
    # EXCEL-EXPORT
    # ==================================================================

    def export_excel(self):
        """Exportiert die aktuelle Haltungstabelle in eine XLSM-Datei."""
        data_frame = pd.DataFrame(columns=[
            "Haltungsname",
            "Länge",
            "Dimension",
            "Entwässerungssystem",
            "Anzahl GALs",
            "Kosten Reinigung",
            "Kosten Reinigung TV",
            "Kosten TV",
            "Kosten GAL",
            "Kosten Panoramo",
        ])

        for row in range(self.dialog.Liste_Haltungen.rowCount()):
            data_row = []
            for column in range(self.dialog.Liste_Haltungen.columnCount()):
                item = self.dialog.Liste_Haltungen.item(row, column)
                data_row.append(item.text() if item else "")
            data_frame.loc[len(data_frame)] = data_row

        template_file_path, _ = QFileDialog.getOpenFileName(
            self.dialog,
            "Excel-Vorlage auswählen",
            "",
            "Excel-Dateien (*.xlsm)",
        )
        if not template_file_path:
            return

        try:
            workbook = openpyxl.load_workbook(
                template_file_path,
                keep_vba=True
            )
            sheet = workbook.active

            for row in dataframe_to_rows(data_frame, index=False, header=True):
                sheet.append(row)

            for row_idx, row_data in enumerate(data_frame.values, 2):
                for col_idx, cell_data in enumerate(row_data):
                    cell = sheet.cell(row=row_idx, column=col_idx + 1)
                    if isinstance(cell_data, (int, float)):
                        cell.number_format = "#,##0.00 [$€-407];[Red]-#,##0.00 [$€-407]"

            export_file_path, _ = QFileDialog.getSaveFileName(
                self.dialog,
                "Excel-Datei speichern",
                "",
                "Excel-Dateien (*.xlsm)",
            )

            if export_file_path:
                workbook.save(export_file_path)
                QMessageBox.information(
                    self.dialog,
                    "Erfolg",
                    f"Datei gespeichert:\n{export_file_path}",
                )

                try:
                    if os.name == "nt":
                        os.startfile(export_file_path)
                    else:
                        subprocess.Popen(["xdg-open", export_file_path])
                except Exception as e:
                    print(f"Fehler beim Öffnen der Datei: {e}")

        except Exception as e:
            QMessageBox.critical(
                self.dialog,
                "Fehler beim Export",
                f"Fehler: {e}",
            )